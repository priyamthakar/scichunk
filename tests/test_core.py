from pathlib import Path

from docx import Document

from scichunk import SciChunker


def test_process_markdown_sections(tmp_path: Path) -> None:
    paper = tmp_path / "paper.md"
    paper.write_text(
        "Abstract\nThis is a short abstract.\n\n"
        "Introduction\nQuercetin was studied with PCL nanoparticles [12].\n\n"
        "Methods\nNanoprecipitation and DLS were used.",
        encoding="utf-8",
    )

    chunks = SciChunker(research_mode="pharma", max_chunk_tokens=200, target_model="claude").process(paper)

    assert chunks
    assert {chunk.section for chunk in chunks} >= {"Abstract", "Introduction", "Methods"}
    intro_chunk = next(chunk for chunk in chunks if chunk.section == "Introduction")
    assert "[12]" in intro_chunk.references_cited
    assert "quercetin" in [entity.lower() for entity in intro_chunk.entities["drugs"]]
    assert intro_chunk.word_count > 0
    assert intro_chunk.char_count > 0
    assert intro_chunk.target_model == "claude"


def test_process_docx_sections(tmp_path: Path) -> None:
    file_path = tmp_path / "paper.docx"
    document = Document()
    document.add_paragraph("Abstract")
    document.add_paragraph("Curcumin loaded chitosan nanoparticles were prepared.")
    document.add_paragraph("Methods")
    document.add_paragraph("DLS and HPLC were used for characterization.")
    document.save(file_path)

    chunks = SciChunker(research_mode="pharma", max_chunk_tokens=200).process(file_path)

    assert {chunk.section for chunk in chunks} >= {"Abstract", "Methods"}
    all_entities = {key: value for chunk in chunks for key, value in chunk.entities.items()}
    assert "drugs" in all_entities
    assert "techniques" in all_entities


def test_unsupported_file_type(tmp_path: Path) -> None:
    file_path = tmp_path / "paper.csv"
    file_path.write_text("x,y", encoding="utf-8")

    try:
        SciChunker().process(file_path)
    except ValueError as exc:
        assert "Unsupported file type" in str(exc)
    else:
        raise AssertionError("Expected ValueError for unsupported file type")


def test_preserves_leading_text_before_first_heading(tmp_path: Path) -> None:
    paper = tmp_path / "front_matter.md"
    paper.write_text(
        "A short title and author block.\n\n"
        "Abstract\nThe abstract starts here.\n\n"
        "Introduction\nMain body follows.",
        encoding="utf-8",
    )

    chunks = SciChunker(max_chunk_tokens=200).process(paper)

    assert [chunk.section for chunk in chunks][:2] == ["Document", "Abstract"]
    assert "author block" in chunks[0].text


def test_supports_roman_numeral_headings(tmp_path: Path) -> None:
    paper = tmp_path / "roman_headings.md"
    paper.write_text(
        "I. Introduction\nQuercetin appears here.\n\n"
        "II. Methods\nDLS was used.",
        encoding="utf-8",
    )

    chunks = SciChunker(research_mode="pharma", max_chunk_tokens=200).process(paper)

    assert {chunk.section for chunk in chunks} >= {"Introduction", "Methods"}
